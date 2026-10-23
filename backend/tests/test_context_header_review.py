"""
Test S4 — duyệt header hàng loạt (backend/scripts/context_header_review.py): gom nhóm tài liệu trùng nội
dung, chia sẻ header tin cậy giữa các bản trùng, lập thay đổi từ Excel đã duyệt (hàm thuần), vòng khép
kín export → đọc lại Excel. Không đụng MongoDB/ChromaDB production (file tạm).

Chạy: python backend/tests/test_context_header_review.py
"""
import sys, os, json, tempfile
sys.stdout.reconfigure(encoding="utf-8")
sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "scripts")))

import context_header_review as review

H_LAW = "[Luật Doanh nghiệp số 59/2020/QH14]"


def row(filename, header, source, needs_review=False, doc_id=None):
    return {"doc_id": doc_id or filename, "filename": filename, "header": header, "source": source,
            "needs_review": needs_review, "system_header": None, "note": "", "equivalent_to": [], "kind": "legal",
            "chunk_count": 3, "opening": "...", "current_approved": None, "current_suggested": None}


# ══════════════════════════════════════════════════════════════
print("=" * 60)
print("TEST 1: equivalence_groups — trùng chunk + danh sách xác minh tay")
hashes = {"a.pdf": {"1", "2", "3"}, "a_copy.pdf": {"1", "2", "3"}, "b.pdf": {"1", "9", "8", "7"},
          "luatdoanhnghiep.pdf": {"x"}, "59.signed.pdf": {"y"}}
groups = review.equivalence_groups(hashes, [["59.signed.pdf", "luatdoanhnghiep.pdf"], ["khong-co.pdf", "b.pdf"]])
assert sorted(groups) == [["59.signed.pdf", "luatdoanhnghiep.pdf"], ["a.pdf", "a_copy.pdf"]], groups
print("  ✅ PASS — b.pdf (trùng 1/4) không bị gom; nhóm tay bỏ qua tên không tồn tại\n")

# ══════════════════════════════════════════════════════════════
print("=" * 60)
print("TEST 2: propagate_twins — header tin cậy sang bản trùng, ghi rõ nguồn")
rows = [row("luatdoanhnghiep.pdf", H_LAW, "rule"), row("59.signed.pdf", "[Điều lệ]", "slm", True),
        row("45.signed.pdf", None, None, True), row("Boluatlaodong2019.pdf", None, None, True)]
review.propagate_twins(rows, [["59.signed.pdf", "luatdoanhnghiep.pdf"], ["45.signed.pdf", "Boluatlaodong2019.pdf"]])
by = {r["filename"]: r for r in rows}
assert by["59.signed.pdf"]["header"] == H_LAW and by["59.signed.pdf"]["source"] == "twin"
assert by["59.signed.pdf"]["system_header"] == "[Điều lệ]" and by["59.signed.pdf"]["needs_review"] is True
assert by["luatdoanhnghiep.pdf"]["header"] == H_LAW and by["luatdoanhnghiep.pdf"]["source"] == "rule"
assert by["45.signed.pdf"]["header"] is None, "nhóm không có header tin cậy → không bịa"
assert by["45.signed.pdf"]["equivalent_to"] == ["Boluatlaodong2019.pdf"], "vẫn ghi thông tin trùng để người duyệt biết"

poc_rows = [row("x-BCTCR.pdf", "[BCTC riêng 2019]", "poc"), row("y-BCTCR.pdf", "[BCTC 2019 nhầm]", "rule", True)]
review.propagate_twins(poc_rows, [["x-BCTCR.pdf", "y-BCTCR.pdf"]])
assert poc_rows[1]["header"] == "[BCTC riêng 2019]", "header PoC đã duyệt là nguồn tin cậy nhất"
print("  ✅ PASS\n")

# ══════════════════════════════════════════════════════════════
print("=" * 60)
print("TEST 3: plan_review_changes — approved / no_header / pending + lỗi dữ liệu")
rows = [row("a.pdf", "[gợi ý A]", "rule", doc_id="d1"), row("b.pdf", "[gợi ý B]", "slm", True, doc_id="d2"),
        row("c.pdf", None, None, True, doc_id="d3"), row("d.pdf", "[gợi ý D]", "rule", doc_id="d4")]
current = {d: {"_id": d} for d in ("d1", "d2", "d3", "d4")}
reviews = {
    "d1": {"status": "approved", "final_header": "[gợi ý A]", "review_note": ""},
    "d2": {"status": "approved", "final_header": "  Sửa lại   tiêu đề B ", "review_note": "SLM sai loại"},
    "d3": {"status": "no_header", "final_header": "", "review_note": ""},
    "d4": {"status": "pending", "final_header": "[gợi ý D]", "review_note": ""},
}
changes = {c["doc_id"]: c for c in review.plan_review_changes(rows, reviews, current)}
assert set(changes) == {"d1", "d2", "d3"}, "pending không ghi gì"
u1 = changes["d1"]["update"]
assert (u1["context_header"], u1["context_header_suggested"], u1["context_header_source"], u1["context_header_version"]) \
    == ("[gợi ý A]", "[gợi ý A]", "rule", 1)
u2 = changes["d2"]["update"]
assert u2["context_header"] == "[Sửa lại tiêu đề B]" and u2["context_header_source"] == "user"
assert u2["context_header_suggested"] == "[gợi ý B]" and u2["context_header_review_note"] == "SLM sai loại"
u3 = changes["d3"]["update"]
assert (u3["context_header"], u3["context_header_suggested"], u3["context_header_version"]) == (None, None, 0)

for bad_reviews, message in (
    ({"d1": {"status": "duyệt", "final_header": "[x]", "review_note": ""}}, "Trạng thái không hợp lệ"),
    ({"zzz": {"status": "approved", "final_header": "[x]", "review_note": ""}}, "không có trong danh sách"),
    ({"d1": {"status": "approved", "final_header": "  ", "review_note": ""}}, "Header chốt' rỗng"),
    ({"d1": {"status": "approved", "final_header": "x" * 300, "review_note": ""}}, "tối đa"),
):
    try:
        review.plan_review_changes(rows, bad_reviews, current)
        raise AssertionError(f"phải lỗi: {message}")
    except ValueError as exc:
        assert message in str(exc), exc
try:
    review.plan_review_changes(rows, {"d1": reviews["d1"]}, {})
    raise AssertionError("tài liệu đã bị xóa khỏi MongoDB phải báo lỗi")
except ValueError as exc:
    assert "không còn trong MongoDB" in str(exc)
print("  ✅ PASS\n")

# ══════════════════════════════════════════════════════════════
print("=" * 60)
print("TEST 4: export → sửa Excel → read_review_sheet (vòng khép kín, file tạm)")
from openpyxl import load_workbook

tmp = tempfile.mkdtemp(prefix="ctx_review_")
review.SUGGESTIONS_PATH = os.path.join(tmp, "s.json")
review.SHEET_PATH = os.path.join(tmp, "r.xlsx")
sample = [row("a.pdf", "[gợi ý A]", "poc", doc_id="d1"), row("b.pdf", None, None, True, doc_id="d2")]
sample[1]["equivalent_to"] = ["c.pdf"]
with open(review.SUGGESTIONS_PATH, "w", encoding="utf-8") as f:
    json.dump({"rows": sample}, f, ensure_ascii=False)
review.cmd_export(force=False)

wb = load_workbook(review.SHEET_PATH)
ws = wb[review.REVIEW_SHEET]
col = {k: i + 1 for i, (_, _, k) in enumerate(review.COLUMNS)}
assert ws.cell(2, col["status"]).value == "approved", "dòng PoC điền sẵn approved"
assert ws.cell(3, col["status"]).value == "pending" and ws.cell(3, col["final_header"]).value in ("", None)
assert "c.pdf" in ws.cell(3, col["info"]).value
assert [d.formula1 for d in ws.data_validations.dataValidation] == ['"approved,no_header,pending"']
ws.cell(3, col["status"]).value = "approved"
ws.cell(3, col["final_header"]).value = "Quy trình B — Đơn vị X, 2023"
ws.cell(3, col["header"]).value = "SỬA Ở CỘT CHỈ ĐỌC — PHẢI BỊ BỎ QUA"
wb.save(review.SHEET_PATH)

read = review.read_review_sheet(review.SHEET_PATH)
assert read["d1"] == {"status": "approved", "final_header": "[gợi ý A]", "review_note": ""}
assert read["d2"]["final_header"] == "Quy trình B — Đơn vị X, 2023"
changes = review.plan_review_changes(sample, read, {"d1": {"_id": "d1"}, "d2": {"_id": "d2"}})
assert {c["doc_id"]: c["after"] for c in changes} == {"d1": "[gợi ý A]", "d2": "[Quy trình B — Đơn vị X, 2023]"}
try:
    review.cmd_export(force=False)
    raise AssertionError("không được ghi đè file Excel đang duyệt dở")
except SystemExit:
    pass
print("  ✅ PASS\n")

print("🎉 TẤT CẢ TEST DUYỆT HEADER HÀNG LOẠT PASS!")
